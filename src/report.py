"""
Builds an Excel reconciliation report from reconcile.py's output: a Summary
sheet with counts by status, a Trade Detail sheet with every trade
color-coded by status, and a Duplicates sheet when any exist.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

STATUS_COLORS = {
    "matched": "C6EFCE",           # green
    "missing_external": "FFC7CE",  # red
    "missing_internal": "FFC7CE",  # red
    "break": "FFEB9C",             # amber
}
DUPLICATE_COLOR = "FFD966"


def _write_dataframe(ws, df: pd.DataFrame, status_col: str | None = None, fixed_color: str | None = None):
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)

    status_idx = list(df.columns).index(status_col) if status_col else None
    for row in dataframe_to_rows(df, index=False, header=False):
        ws.append(row)
        color = fixed_color
        if status_idx is not None:
            color = STATUS_COLORS.get(row[status_idx])
        if color:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")


def write_report(results_df: pd.DataFrame, duplicates_df: pd.DataFrame, output_path: str) -> None:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    counts = results_df["status"].value_counts().to_dict()
    if not duplicates_df.empty:
        counts["duplicate_entry"] = int(len(duplicates_df))
    summary_df = pd.DataFrame(sorted(counts.items()), columns=["status", "count"])
    _write_dataframe(summary_ws, summary_df)

    detail_ws = wb.create_sheet("Trade Detail")
    _write_dataframe(detail_ws, results_df, status_col="status")

    if not duplicates_df.empty:
        dup_ws = wb.create_sheet("Duplicates")
        _write_dataframe(dup_ws, duplicates_df, fixed_color=DUPLICATE_COLOR)

    wb.save(output_path)
